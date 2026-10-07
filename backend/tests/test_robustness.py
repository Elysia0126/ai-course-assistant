"""Index versioning, upload hardening and passage lookup."""

import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import update

from app.models import Document
from app.services.parsing import parse_document
from tests import factories
from tests.conftest import upload


def test_changing_embedding_model_requires_reindex(client: TestClient, ready_course: dict, app) -> None:
    cid = ready_course["id"]
    docs = client.get(f"/api/courses/{cid}/documents").json()
    assert all(d["embedding_signature"] == "hash:feature-hash-384:384" and not d["needs_reindex"] for d in docs)

    # Simulate documents indexed by another model (e.g. after switching EMBEDDING_MODEL).
    with app.state.engine.begin() as conn:
        conn.execute(update(Document).where(Document.course_id == cid).values(embedding_signature="openai:x:384"))

    assert all(d["needs_reindex"] for d in client.get(f"/api/courses/{cid}/documents").json())
    stale = client.post(f"/api/courses/{cid}/search", json={"query": "learning rate"})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "index_outdated"
    assert client.post(f"/api/courses/{cid}/chat/stream", json={"question": "learning rate?"}).status_code == 409
    topic_quiz = client.post(f"/api/courses/{cid}/quizzes", json={"topic": "learning rate"})
    assert topic_quiz.status_code == 409

    reindexed = client.post(f"/api/courses/{cid}/reindex")
    assert reindexed.status_code == 202 and len(reindexed.json()) == 3
    docs = client.get(f"/api/courses/{cid}/documents").json()
    assert {d["status"] for d in docs} == {"ready"} and not any(d["needs_reindex"] for d in docs)
    assert client.post(f"/api/courses/{cid}/search", json={"query": "learning rate"}).status_code == 200


def test_zip_bomb_office_file_is_rejected(client: TestClient, course: dict, settings, monkeypatch) -> None:
    monkeypatch.setattr(settings, "max_unzipped_mb", 1)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", b"0" * (3 * 1024 * 1024))  # 3 MB → a few KB compressed
    data = buffer.getvalue()
    assert len(data) < 50_000
    response = upload(client, course["id"], ("bomb.pptx", data, "application/octet-stream"))
    assert response.status_code == 413
    assert "expands to more than 1 MB" in response.json()["error"]["message"]


def test_page_limit_marks_document_failed(client: TestClient, course: dict, settings, monkeypatch) -> None:
    monkeypatch.setattr(settings, "max_pages", 2)
    response = upload(client, course["id"], ("long.pdf", factories.make_pdf(), "application/pdf"))
    document = client.get(f"/api/documents/{response.json()['documents'][0]['id']}").json()
    assert document["status"] == "failed"
    assert "the limit is 2" in document["error_message"]


def test_text_limit(tmp_path: Path) -> None:
    path = tmp_path / "big.md"
    path.write_text("word " * 5000, encoding="utf-8")
    with pytest.raises(Exception, match="more than 1,000 characters"):
        parse_document(path, "md", max_chars=1000)


def test_grouped_shapes_in_slides_are_parsed(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Grouped diagram"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Text hidden inside a grouped shape"
    path = tmp_path / "group.pptx"
    prs.save(path)
    parsed = parse_document(path, "pptx")
    assert "Text hidden inside a grouped shape" in parsed.sections[0].text


def test_chunk_detail_returns_full_passage(client: TestClient, ready_course: dict) -> None:
    hit = client.post(f"/api/courses/{ready_course['id']}/search", json={"query": "warmup"}).json()["results"][0]
    chunk = client.get(f"/api/chunks/{hit['chunk_id']}").json()
    assert chunk["filename"] == hit["filename"] and chunk["page_number"] == hit["page_number"]
    assert len(chunk["content"]) >= len(hit["snippet"].rstrip("…"))
    assert client.get("/api/chunks/does-not-exist").status_code == 404
