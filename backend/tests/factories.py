"""Build small but realistic course files on the fly for tests."""

import io

GRADIENT_PAGES = [
    (
        "Gradient Descent",
        "Gradient descent is an iterative optimization algorithm that updates parameters in the direction of the "
        "negative gradient. The learning rate is a hyperparameter that controls the size of each update step.",
    ),
    (
        "Learning Rate Schedules",
        "If the learning rate is too large the loss diverges, and if it is too small training converges slowly. "
        "Warmup is a technique that starts training with a small learning rate and increases it gradually. "
        "Cosine annealing lowers the learning rate smoothly following a cosine curve.",
    ),
    (
        "Adaptive Optimizers",
        "Momentum is a method that accumulates an exponentially decaying average of past gradients. Adam is an "
        "adaptive optimizer that combines momentum with per-parameter learning rates estimated from the first and "
        "second moments of the gradients.",
    ),
]


def make_pdf(pages: list[tuple[str, str]] = GRADIENT_PAGES) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    story = []
    for i, (title, body) in enumerate(pages):
        story += [Paragraph(title, styles["Heading1"]), Paragraph(body, styles["BodyText"])]
        if i < len(pages) - 1:
            story.append(PageBreak())
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=LETTER).build(story)
    return buffer.getvalue()


def make_blank_pdf() -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slides = [
        (
            "Backpropagation",
            ["Backpropagation is an algorithm that computes gradients with the chain rule"],
            "It runs backwards through the network after the forward pass.",
        ),
        (
            "Activation Functions",
            ["ReLU is an activation function defined as max(0, x)", "Sigmoid outputs (0, 1)"],
            "ReLU avoids saturation for positive inputs.",
        ),
    ]
    for title, bullets, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
        slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Dropout", level=1)
    doc.add_paragraph("Dropout is a regularization technique that randomly zeroes activations during training.")
    doc.add_heading("Early Stopping", level=1)
    doc.add_paragraph("Early stopping halts training when validation loss stops improving.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


MARKDOWN = b"""# Regularization

## Overfitting
Overfitting is when a model learns noise in the training data and performs poorly on unseen data.

## L2 Regularization
L2 regularization adds a penalty proportional to the squared magnitude of the weights to the loss.

## Dropout
Dropout is a regularization technique that randomly sets a fraction of activations to zero during training.
"""
