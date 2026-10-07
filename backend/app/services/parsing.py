"""Turn uploaded files into location-aware text sections.

Each ``ParsedSection`` remembers where it came from (PDF page, slide number, or heading) so that
answers can cite "Lecture 3.pdf, p. 4" instead of just a file name.
"""

import logging
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.errors import DocumentParsingError
from app.services.text_utils import clean_whitespace, reflow_lines

logger = logging.getLogger(__name__)

SUPPORTED_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
}


@dataclass
class ParsedSection:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass
class ParsedDocument:
    sections: list[ParsedSection] = field(default_factory=list)
    page_count: int | None = None

    @property
    def total_chars(self) -> int:
        return sum(len(s.text) for s in self.sections)


def file_type_for(filename: str) -> str:
    """Normalised type key ('pdf', 'pptx', 'docx', 'md', 'txt') for a filename."""
    suffix = Path(filename).suffix.lower()
    return "md" if suffix == ".markdown" else suffix.lstrip(".")


def parse_document(
    path: Path, file_type: str, *, max_pages: int | None = None, max_chars: int | None = None
) -> ParsedDocument:
    parsers = {"pdf": parse_pdf, "pptx": parse_pptx, "docx": parse_docx, "md": parse_markdown, "txt": parse_text}
    parser = parsers.get(file_type)
    if parser is None:
        raise DocumentParsingError(f"No parser available for .{file_type} files.")
    try:
        parsed = parser(path, max_pages=max_pages) if file_type in {"pdf", "pptx"} else parser(path)
    except DocumentParsingError:
        raise
    except Exception as exc:  # corrupt files surface as library-specific exceptions
        logger.warning("Failed to parse %s: %s", path.name, exc)
        raise DocumentParsingError(
            f"Could not read this .{file_type} file — it may be corrupted or encrypted."
        ) from exc

    parsed.sections = [s for s in parsed.sections if s.text.strip()]
    if not parsed.sections:
        hint = " It may be a scanned PDF (images only); OCR is not supported yet." if file_type == "pdf" else ""
        raise DocumentParsingError(f"No extractable text was found in this file.{hint}")
    if max_chars and parsed.total_chars > max_chars:
        raise DocumentParsingError(
            f"This file contains more than {max_chars:,} characters of text. Split it into smaller documents."
        )
    return parsed


def _check_page_limit(count: int, max_pages: int | None, unit: str) -> None:
    if max_pages and count > max_pages:
        raise DocumentParsingError(f"This file has {count} {unit}; the limit is {max_pages}. Split it and retry.")


def parse_pdf(path: Path, max_pages: int | None = None) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentParsingError("This PDF is password-protected.") from exc
    _check_page_limit(len(reader.pages), max_pages, "pages")

    sections: list[ParsedSection] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # one broken page should not sink the whole document
            logger.warning("Skipping unreadable page %s of %s: %s", index, path.name, exc)
            continue
        text = reflow_lines(clean_whitespace(text))
        if text:
            sections.append(ParsedSection(text=text, page_number=index, section=_first_line_heading(text)))
    return ParsedDocument(sections=sections, page_count=len(reader.pages))


def _shape_lines(shapes: Iterable[Any], skip: Any) -> Iterator[str]:
    """Text from slide shapes, recursing into grouped shapes (common in real lecture decks)."""
    for shape in shapes:
        if shape == skip:
            continue
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # MSO_SHAPE_TYPE.GROUP
            yield from _shape_lines(shape.shapes, skip)
            continue
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                line = "".join(run.text for run in paragraph.runs).strip()
                if line:
                    yield ("  " * paragraph.level) + "- " + line if paragraph.level else line
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text.strip() for cell in row.cells)


def parse_pptx(path: Path, max_pages: int | None = None) -> ParsedDocument:
    from pptx import Presentation

    presentation = Presentation(str(path))
    _check_page_limit(len(presentation.slides), max_pages, "slides")
    sections: list[ParsedSection] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip() or None

        lines = list(_shape_lines(slide.shapes, skip=slide.shapes.title))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")

        body = "\n".join(lines)
        text = f"{title}\n{body}" if title else body
        text = clean_whitespace(text)
        if text:
            sections.append(ParsedSection(text=text, page_number=index, section=title))
    return ParsedDocument(sections=sections, page_count=len(presentation.slides))


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    sections: list[ParsedSection] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = clean_whitespace("\n".join(buffer))
        if text:
            sections.append(ParsedSection(text=text, section=heading))
        buffer.clear()

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            heading = text[:300]
            buffer.append(text)
        else:
            buffer.append(text)
    flush()

    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = clean_whitespace("\n".join(rows))
        if text:
            sections.append(ParsedSection(text=text, section="Table"))
    return ParsedDocument(sections=sections)


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def parse_markdown(path: Path) -> ParsedDocument:
    raw = _read_text(path)
    sections: list[ParsedSection] = []
    stack: list[str] = []
    buffer: list[str] = []
    in_code = False

    def flush() -> None:
        text = clean_whitespace("\n".join(buffer))
        if text:
            # Skip the H1 (usually the document title) and keep at most the two deepest headings.
            path = stack[1:] if len(stack) > 1 else stack
            sections.append(ParsedSection(text=text, section=" › ".join(path[-2:])[:300] or None))
        buffer.clear()

    for line in raw.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        match = None if in_code else _MD_HEADING.match(line)
        if match:
            flush()
            level = len(match.group(1))
            stack[:] = [*stack[: level - 1], match.group(2).strip()]
            buffer.append(match.group(2).strip())
        else:
            buffer.append(line)
    flush()
    return ParsedDocument(sections=sections)


def parse_text(path: Path) -> ParsedDocument:
    text = clean_whitespace(_read_text(path))
    return ParsedDocument(sections=[ParsedSection(text=text)] if text else [])


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _first_line_heading(text: str) -> str | None:
    """Use a short first line (typical slide/page title) as the section label."""
    first = text.split("\n", 1)[0].strip()
    if 3 <= len(first) <= 90 and not first.endswith((".", ",", ";")):
        return first
    return None
