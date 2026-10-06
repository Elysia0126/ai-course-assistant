"""Generate the demo course materials in ../sample_data (original content written for this project).

    python scripts/make_sample_data.py

Produces a 6-page PDF lecture, an 8-slide PPTX deck with speaker notes, and a Markdown study guide —
one file per supported "shape" of course material.
"""

from pathlib import Path

OUT_DIR = Path(__file__).resolve().parents[2] / "sample_data"

PDF_PAGES: list[tuple[str, list[str]]] = [
    (
        "Lecture 3: Optimization with Gradient Descent",
        [
            "Machine learning models are trained by solving an optimization problem: we search for the "
            "parameters that make the model's predictions match the training data as closely as possible.",
            "A loss function is a function that measures how far a model's predictions are from the true targets. "
            "Common examples are the squared error for regression and the cross-entropy loss for classification.",
            "Empirical risk is the average loss over all examples in the training set. Training a model means "
            "minimizing the empirical risk with respect to the model parameters, usually written as theta.",
            "For most models there is no closed-form solution, so we use iterative optimization methods. "
            "This lecture covers gradient descent and the variants that are used to train modern neural networks.",
        ],
    ),
    (
        "The Gradient Descent Algorithm",
        [
            "The gradient is the vector of partial derivatives of the loss with respect to each parameter. "
            "It points in the direction of steepest increase of the loss.",
            "Gradient descent is an iterative optimization algorithm that updates the parameters in the direction "
            "of the negative gradient. Each iteration applies the update rule: "
            "theta_new = theta - eta * grad L(theta).",
            "The learning rate is a hyperparameter that controls the size of each update step. It is written as "
            "eta and is usually between 0.0001 and 0.1.",
            "The algorithm repeats the update until the loss stops decreasing, a maximum number of iterations is "
            "reached, or the norm of the gradient falls below a small tolerance.",
        ],
    ),
    (
        "Choosing the Learning Rate",
        [
            "If the learning rate is too large, the updates overshoot the minimum: the loss oscillates or even "
            "diverges to infinity. If the learning rate is too small, training converges very slowly and may "
            "stall on plateaus.",
            "A learning rate schedule changes the learning rate during training. Popular schedules include step "
            "decay, exponential decay and cosine annealing, which smoothly lowers the rate following a cosine curve.",
            "Warmup is a technique that starts training with a small learning rate and increases it gradually over "
            "the first few hundred steps. Warmup stabilizes the early phase of training for large models.",
            "A practical way to pick an initial value is a learning rate range test: train briefly while increasing "
            "the rate exponentially and choose a value slightly below the point where the loss starts to explode.",
        ],
    ),
    (
        "Stochastic and Mini-batch Gradient Descent",
        [
            "Batch gradient descent computes the exact gradient over the entire training set, which is expensive "
            "when the dataset contains millions of examples.",
            "Stochastic gradient descent (SGD) estimates the gradient using a single randomly chosen training "
            "example. Each step is cheap but noisy.",
            "Mini-batch gradient descent computes the gradient on a small random subset of examples called a batch. "
            "Typical batch sizes range from 32 to 256 and make good use of GPU parallelism.",
            "An epoch is one full pass over the training dataset. The noise in stochastic updates can help the "
            "optimizer escape saddle points and shallow local minima.",
        ],
    ),
    (
        "Momentum and Adaptive Methods",
        [
            "Momentum is a method that accumulates an exponentially decaying average of past gradients to "
            "accelerate descent. It dampens oscillations in steep directions and speeds up progress along "
            "shallow, consistent directions.",
            "RMSProp divides the learning rate for each parameter by a running average of recent squared gradients, "
            "so parameters with large gradients take smaller steps.",
            "Adam is an adaptive optimizer that combines momentum with per-parameter learning rates based on "
            "estimates of the first and second moments of the gradients. Its default hyperparameters are "
            "beta1 = 0.9, beta2 = 0.999 and epsilon = 1e-8.",
            "AdamW decouples weight decay from the adaptive gradient update and is the default optimizer for "
            "training transformer models.",
        ],
    ),
    (
        "Convergence and Practical Tips",
        [
            "For convex loss functions, gradient descent with a suitable learning rate is guaranteed to converge to "
            "the global minimum. Neural network losses are non-convex and contain many local minima and saddle points.",
            "Feature scaling speeds up convergence: standardizing each input feature to zero mean and unit variance "
            "makes the loss surface more round, so gradient steps point closer to the minimum.",
            "Gradient clipping is a technique that rescales gradients whose norm exceeds a threshold. It prevents "
            "exploding gradients in recurrent networks and very deep models.",
            "Always monitor both training loss and validation loss. A training loss that decreases while the "
            "validation loss increases is a sign of overfitting, which is covered in the regularization notes.",
        ],
    ),
]

SLIDES: list[tuple[str, list[str], str]] = [
    (
        "Lecture 5: Neural Networks and Backpropagation",
        ["Machine Learning Foundations", "From perceptrons to deep networks"],
        "Today we build up from a single neuron to multi-layer networks and derive how they are trained.",
    ),
    (
        "The Perceptron",
        [
            "A perceptron is a linear classifier that outputs 1 if the weighted sum of its inputs exceeds a threshold",
            "Weights w and bias b are learned from labelled data",
            "Can only separate classes with a straight line (hyperplane)",
        ],
        "The XOR problem shows the limitation: no single line separates the XOR classes, which motivated "
        "multi-layer networks.",
    ),
    (
        "Activation Functions",
        [
            "Sigmoid squashes values into the range (0, 1)",
            "Tanh squashes values into the range (-1, 1) and is zero-centred",
            "ReLU is an activation function defined as max(0, x)",
            "Non-linear activations let networks represent non-linear functions",
        ],
        "Without a non-linear activation, stacking layers collapses into a single linear transformation. "
        "ReLU is the default choice for hidden layers because it is cheap and does not saturate for positive inputs.",
    ),
    (
        "Multi-layer Perceptrons",
        [
            "Input layer, one or more hidden layers, output layer",
            "Each layer computes h = activation(W x + b)",
            "Universal approximation theorem: one hidden layer with enough units can approximate "
            "any continuous function",
            "Depth lets networks reuse features and learn hierarchies",
        ],
        "In practice deeper networks with fewer units per layer generalize better than very wide shallow ones.",
    ),
    (
        "The Forward Pass",
        [
            "Compute activations layer by layer from input to output",
            "Apply softmax at the output to get class probabilities",
            "Compare predictions with labels using the cross-entropy loss",
        ],
        "The forward pass caches intermediate activations because backpropagation needs them.",
    ),
    (
        "Backpropagation",
        [
            "Backpropagation is an algorithm that computes the gradient of the loss with respect to every weight "
            "by applying the chain rule backwards through the network",
            "Cost is roughly the same as one forward pass",
            "Gradients are then used by an optimizer such as SGD or Adam",
        ],
        "Backpropagation is reverse-mode automatic differentiation. Frameworks like PyTorch build a computation "
        "graph during the forward pass and traverse it in reverse to accumulate gradients.",
    ),
    (
        "Vanishing and Exploding Gradients",
        [
            "The vanishing gradient problem occurs when gradients shrink exponentially as they are propagated back "
            "through many layers",
            "Exploding gradients grow exponentially and make training unstable",
            "Fixes: ReLU activations, He or Xavier initialization, batch normalization, residual connections",
        ],
        "Residual connections add the input of a block to its output, giving gradients a short path to early layers.",
    ),
    (
        "Training Checklist",
        [],
        "Use this table as a debugging guide when a network does not train.",
    ),
]

CHECKLIST_TABLE = [
    ("Symptom", "Likely cause", "Fix"),
    ("Loss is NaN", "Learning rate too high", "Lower the learning rate, clip gradients"),
    ("Loss does not decrease", "Bug or learning rate too low", "Overfit one batch first"),
    ("Train good, validation bad", "Overfitting", "Regularization, more data"),
]

MARKDOWN_NOTES = """# Overfitting and Regularization — Study Guide

These notes accompany Lecture 6 of Machine Learning Foundations.

## Bias and Variance

The bias-variance trade-off describes how prediction error splits into two sources. Bias is the error caused by
overly simple assumptions in the model; high-bias models underfit. Variance is the error caused by sensitivity to
small fluctuations in the training set; high-variance models overfit.

## Overfitting

Overfitting is when a model learns noise and accidental patterns in the training data and therefore performs poorly
on unseen data. The classic symptom is a training loss that keeps falling while the validation loss starts to rise.
Underfitting is the opposite: the model is too simple to capture the structure of the data.

## L2 Regularization (Weight Decay)

L2 regularization adds a penalty proportional to the squared magnitude of the weights to the loss:
`L_total = L_data + lambda * ||w||^2`. Larger values of lambda push weights towards zero and produce smoother
models. In neural networks this is usually implemented as weight decay.

## L1 Regularization

L1 regularization adds a penalty proportional to the absolute value of the weights. Unlike L2, it drives many
weights exactly to zero, which produces sparse models and acts as a form of feature selection.

## Dropout

Dropout is a regularization technique that randomly sets a fraction of activations to zero during training. A
typical dropout rate is 0.5 for fully connected layers. At test time all units are used and activations are scaled
so their expected value matches training.

## Early Stopping

Early stopping is a strategy that halts training when the validation loss stops improving for a number of epochs,
called the patience. The weights from the best validation epoch are restored at the end.

## Cross-Validation

k-fold cross-validation splits the data into k folds, trains on k-1 folds and evaluates on the remaining fold,
rotating k times. The average score is a more reliable estimate of generalization than a single split, and it is
the standard way to choose hyperparameters such as lambda.

## Data Augmentation

Data augmentation creates new training examples by applying label-preserving transformations, such as flips,
crops and colour jitter for images. More varied data reduces variance without changing the model.
"""


def make_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    for i, (title, paragraphs) in enumerate(PDF_PAGES):
        story.append(Paragraph(title, styles["Title" if i == 0 else "Heading1"]))
        story.append(Spacer(1, 12))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, styles["BodyText"]))
            story.append(Spacer(1, 8))
        if i < len(PDF_PAGES) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=LETTER, title="Lecture 3 - Gradient Descent").build(story)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for i, (title, bullets, notes) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        if i == 0:
            slide.placeholders[1].text = "\n".join(bullets)
        elif bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet
        else:
            # Remove the empty body placeholder and add a table instead.
            slide.shapes._spTree.remove(slide.placeholders[1]._element)
            rows, cols = len(CHECKLIST_TABLE), len(CHECKLIST_TABLE[0])
            table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.6), Inches(9), Inches(3)).table
            for r, row in enumerate(CHECKLIST_TABLE):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    cell.text_frame.paragraphs[0].font.size = Pt(16)
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_pdf(OUT_DIR / "Lecture03_Gradient_Descent.pdf")
    make_pptx(OUT_DIR / "Lecture05_Neural_Networks.pptx")
    (OUT_DIR / "Regularization_Study_Guide.md").write_text(MARKDOWN_NOTES, encoding="utf-8")
    for file in sorted(OUT_DIR.iterdir()):
        print(f"wrote {file.relative_to(OUT_DIR.parent)} ({file.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
